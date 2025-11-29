"""
Document Parser Service

Extracts text from various document formats:
- PDF (using pypdf)
- DOCX (using python-docx)
- TXT, MD (plain text)

Returns standardized Document objects with metadata.
"""

from typing import List, Dict, Any
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from langchain_core.documents import Document


class DocumentParser:
    """
    Parse documents and extract text content.

    Supports: PDF, DOCX, TXT, MD
    Returns: LangChain Document objects with metadata
    """

    SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.txt', '.md'}

    @classmethod
    def is_supported(cls, file_path: str) -> bool:
        """Check if file format is supported"""
        return Path(file_path).suffix.lower() in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def parse(cls, file_path: str) -> List[Document]:
        """
        Parse document and return LangChain Document objects.

        Args:
            file_path: Path to document file

        Returns:
            List of Document objects (usually one per file, but PDFs can be split by page)

        Raises:
            ValueError: If file format not supported
            FileNotFoundError: If file doesn't exist
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        if not cls.is_supported(file_path):
            raise ValueError(
                f"Unsupported file format: {path.suffix}. "
                f"Supported: {', '.join(cls.SUPPORTED_EXTENSIONS)}"
            )

        # Route to appropriate parser
        extension = path.suffix.lower()

        if extension == '.pdf':
            return cls._parse_pdf(path)
        elif extension == '.docx':
            return cls._parse_docx(path)
        elif extension in {'.txt', '.md'}:
            return cls._parse_text(path)

        raise ValueError(f"No parser for {extension}")

    @classmethod
    def _parse_pdf(cls, path: Path) -> List[Document]:
        """
        Parse PDF file.

        Returns one Document per page (allows fine-grained retrieval).
        You can also combine all pages into one Document - your choice!
        """
        try:
            reader = PdfReader(str(path))
            documents = []

            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()

                # Skip empty pages
                if not text.strip():
                    continue

                # Create Document with metadata
                doc = Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "file_name": path.name,
                        "file_type": "pdf",
                        "page": page_num + 1,
                        "total_pages": len(reader.pages)
                    }
                )
                documents.append(doc)

            return documents

        except Exception as e:
            raise ValueError(f"Failed to parse PDF {path.name}: {e}")

    @classmethod
    def _parse_docx(cls, path: Path) -> List[Document]:
        """
        Parse DOCX file.

        Returns single Document with all content.
        """
        try:
            doc = DocxDocument(str(path))

            # Extract all paragraphs
            text = "\n".join([para.text for para in doc.paragraphs])

            if not text.strip():
                raise ValueError(f"DOCX file is empty: {path.name}")

            return [
                Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "file_name": path.name,
                        "file_type": "docx",
                        "paragraphs": len(doc.paragraphs)
                    }
                )
            ]

        except Exception as e:
            raise ValueError(f"Failed to parse DOCX {path.name}: {e}")

    @classmethod
    def _parse_text(cls, path: Path) -> List[Document]:
        """
        Parse plain text file (TXT, MD).

        Returns single Document with all content.
        """
        try:
            with open(path, 'r', encoding='utf-8') as f:
                text = f.read()

            if not text.strip():
                raise ValueError(f"Text file is empty: {path.name}")

            return [
                Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "file_name": path.name,
                        "file_type": path.suffix[1:]  # Remove the dot
                    }
                )
            ]

        except UnicodeDecodeError:
            # Try different encoding
            try:
                with open(path, 'r', encoding='latin-1') as f:
                    text = f.read()
                return [
                    Document(
                        page_content=text,
                        metadata={
                            "source": str(path),
                            "file_name": path.name,
                            "file_type": path.suffix[1:],
                            "encoding": "latin-1"
                        }
                    )
                ]
            except Exception as e:
                raise ValueError(f"Failed to parse text file {path.name}: {e}")

        except Exception as e:
            raise ValueError(f"Failed to parse text file {path.name}: {e}")


# =============================================================================
# Test Parser
# =============================================================================
if __name__ == "__main__":
    print("=" * 70)
    print("Document Parser Test")
    print("=" * 70)

    # Test with sample text file
    test_content = """
    RAG (Retrieval-Augmented Generation) is a technique for enhancing LLM responses.

    It works by:
    1. Retrieving relevant documents from a knowledge base
    2. Passing those documents as context to the LLM
    3. Generating an answer based on the retrieved context

    This improves accuracy and allows LLMs to access specific knowledge.
    """

    # Create sample file
    test_file = Path("../data/test_document.txt")
    test_file.write_text(test_content)

    # Test parser
    print(f"\nParsing: {test_file.name}")
    print("-" * 70)

    documents = DocumentParser.parse(str(test_file))

    for doc in documents:
        print(f"Content: {doc.page_content[:100]}...")
        print(f"Metadata: {doc.metadata}")
        print()

    # Cleanup
    test_file.unlink()

    print("=" * 70)
    print("[OK] Parser working!")
    print("=" * 70)
    print("\nNext: Test with real PDF/DOCX files")
    print("Place sample files in data/documents/ folder")
